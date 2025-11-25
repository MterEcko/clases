#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Cursos - Documentos PDF, Excel y PowerPoint
Para cursos de 3 meses (24 sesiones) en Tala, Jalisco, México
"""

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
import os


class GeneradorCurso:
    def __init__(self, curso_info):
        self.info = curso_info
        self.ruta_base = curso_info['ruta']

    def generar_todo(self):
        """Genera todos los documentos del curso"""
        os.makedirs(self.ruta_base, exist_ok=True)

        print(f"Generando curso: {self.info['nombre']} - {self.info['grupo_edad']}")
        self.generar_pdf_plan()
        self.generar_excel_materiales()
        self.generar_powerpoints()
        print(f"✓ Curso generado en: {self.ruta_base}\n")

    def generar_pdf_plan(self):
        """Genera el PDF con el plan completo del curso"""
        archivo = os.path.join(self.ruta_base, "Plan_Curso.pdf")
        doc = SimpleDocTemplate(archivo, pagesize=letter,
                               leftMargin=0.75*inch, rightMargin=0.75*inch,
                               topMargin=0.75*inch, bottomMargin=0.75*inch)

        story = []
        styles = getSampleStyleSheet()

        # Estilos personalizados
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1a5490'),
            spaceAfter=30,
            alignment=TA_CENTER
        )

        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#2c5f8d'),
            spaceAfter=12,
            spaceBefore=12
        )

        # Título
        story.append(Paragraph(self.info['nombre'], title_style))
        story.append(Paragraph(f"Grupo de edad: {self.info['grupo_edad']}", styles['Normal']))
        story.append(Spacer(1, 0.3*inch))

        # Información general
        story.append(Paragraph("INFORMACIÓN GENERAL", heading_style))
        info_general = f"""
        <b>Duración:</b> 3 meses (24 sesiones)<br/>
        <b>Frecuencia:</b> 2 días por semana<br/>
        <b>Duración por sesión:</b> {self.info['duracion_sesion']}<br/>
        <b>Nivel:</b> {self.info['nivel']}<br/>
        <b>Precio del curso:</b> ${self.info['precio_curso']:,.2f} MXN<br/>
        <b>Costo de materiales:</b> ${self.info['costo_materiales']:,.2f} MXN<br/>
        <b>Total:</b> ${self.info['precio_curso'] + self.info['costo_materiales']:,.2f} MXN
        """
        story.append(Paragraph(info_general, styles['Normal']))
        story.append(Spacer(1, 0.2*inch))

        # Descripción
        story.append(Paragraph("DESCRIPCIÓN DEL CURSO", heading_style))
        story.append(Paragraph(self.info['descripcion'], styles['BodyText']))
        story.append(Spacer(1, 0.2*inch))

        # Objetivos de aprendizaje
        story.append(Paragraph("OBJETIVOS DE APRENDIZAJE", heading_style))
        for i, objetivo in enumerate(self.info['objetivos'], 1):
            story.append(Paragraph(f"{i}. {objetivo}", styles['Normal']))
            story.append(Spacer(1, 0.05*inch))
        story.append(Spacer(1, 0.2*inch))

        # Requisitos previos
        story.append(Paragraph("REQUISITOS PREVIOS", heading_style))
        for req in self.info['requisitos']:
            story.append(Paragraph(f"• {req}", styles['Normal']))
        story.append(Spacer(1, 0.2*inch))

        # Índice de módulos
        story.append(PageBreak())
        story.append(Paragraph("ÍNDICE DE MÓDULOS", heading_style))

        for i, modulo in enumerate(self.info['modulos'], 1):
            story.append(Paragraph(f"<b>Módulo {i}: {modulo['nombre']}</b> (Sesiones {modulo['sesiones']})",
                                 styles['Normal']))
            story.append(Paragraph(modulo['descripcion'], styles['Normal']))
            story.append(Spacer(1, 0.15*inch))

        # Plan detallado de sesiones
        story.append(PageBreak())
        story.append(Paragraph("PLAN DETALLADO DE SESIONES", heading_style))

        for i, sesion in enumerate(self.info['sesiones'], 1):
            story.append(Paragraph(f"<b>Sesión {i}: {sesion['titulo']}</b>", styles['Heading3']))
            story.append(Paragraph(f"<b>Duración:</b> {self.info['duracion_sesion']}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))

            story.append(Paragraph("<b>Objetivos:</b>", styles['Normal']))
            for obj in sesion['objetivos']:
                story.append(Paragraph(f"• {obj}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))

            story.append(Paragraph("<b>Contenido:</b>", styles['Normal']))
            for cont in sesion['contenido']:
                story.append(Paragraph(f"• {cont}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))

            story.append(Paragraph(f"<b>Actividades:</b> {sesion['actividades']}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))

            story.append(Paragraph(f"<b>Materiales:</b> {sesion['materiales']}", styles['Normal']))
            story.append(Spacer(1, 0.2*inch))

        # Evaluación
        story.append(PageBreak())
        story.append(Paragraph("SISTEMA DE EVALUACIÓN", heading_style))
        story.append(Paragraph(self.info['evaluacion']['descripcion'], styles['BodyText']))
        story.append(Spacer(1, 0.15*inch))

        story.append(Paragraph("<b>Criterios de evaluación:</b>", styles['Normal']))
        for criterio, peso in self.info['evaluacion']['criterios'].items():
            story.append(Paragraph(f"• {criterio}: {peso}", styles['Normal']))
        story.append(Spacer(1, 0.2*inch))

        # Proyecto final
        story.append(Paragraph("PROYECTO FINAL", heading_style))
        story.append(Paragraph(f"<b>Título:</b> {self.info['proyecto_final']['titulo']}", styles['Normal']))
        story.append(Paragraph(f"<b>Descripción:</b> {self.info['proyecto_final']['descripcion']}",
                             styles['BodyText']))
        story.append(Spacer(1, 0.1*inch))
        story.append(Paragraph("<b>Entregables:</b>", styles['Normal']))
        for entregable in self.info['proyecto_final']['entregables']:
            story.append(Paragraph(f"• {entregable}", styles['Normal']))

        doc.build(story)

    def generar_excel_materiales(self):
        """Genera el Excel con lista de materiales y costos"""
        archivo = os.path.join(self.ruta_base, "Materiales_y_Costos.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Materiales"

        # Estilos
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        title_font = Font(bold=True, size=14)

        # Título
        ws.merge_cells('A1:E1')
        ws['A1'] = f"{self.info['nombre']} - Lista de Materiales"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')

        ws.merge_cells('A2:E2')
        ws['A2'] = f"Grupo de edad: {self.info['grupo_edad']}"
        ws['A2'].alignment = Alignment(horizontal='center')

        # Encabezados
        headers = ['Cantidad', 'Material/Herramienta', 'Descripción', 'Precio Unitario', 'Precio Total']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        # Materiales
        row = 5
        total = 0
        for material in self.info['materiales']:
            ws.cell(row=row, column=1, value=material['cantidad'])
            ws.cell(row=row, column=2, value=material['nombre'])
            ws.cell(row=row, column=3, value=material['descripcion'])
            ws.cell(row=row, column=4, value=material['precio_unitario'])
            ws.cell(row=row, column=4).number_format = '$#,##0.00'

            precio_total = material['cantidad'] * material['precio_unitario']
            ws.cell(row=row, column=5, value=precio_total)
            ws.cell(row=row, column=5).number_format = '$#,##0.00'
            total += precio_total
            row += 1

        # Total
        ws.cell(row=row, column=4, value="TOTAL:")
        ws.cell(row=row, column=4).font = Font(bold=True)
        ws.cell(row=row, column=5, value=total)
        ws.cell(row=row, column=5).font = Font(bold=True)
        ws.cell(row=row, column=5).number_format = '$#,##0.00'

        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 40
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15

        # Resumen de costos
        ws2 = wb.create_sheet(title="Resumen")
        ws2['A1'] = "RESUMEN DE COSTOS"
        ws2['A1'].font = title_font
        ws2['A3'] = "Precio del curso:"
        ws2['B3'] = self.info['precio_curso']
        ws2['B3'].number_format = '$#,##0.00'
        ws2['A4'] = "Costo de materiales:"
        ws2['B4'] = total
        ws2['B4'].number_format = '$#,##0.00'
        ws2['A5'] = "TOTAL:"
        ws2['A5'].font = Font(bold=True)
        ws2['B5'] = self.info['precio_curso'] + total
        ws2['B5'].font = Font(bold=True)
        ws2['B5'].number_format = '$#,##0.00'

        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15

        wb.save(archivo)

    def generar_powerpoints(self):
        """Genera presentaciones PowerPoint para cada módulo"""
        carpeta_ppts = os.path.join(self.ruta_base, "Presentaciones")
        os.makedirs(carpeta_ppts, exist_ok=True)

        for i, modulo in enumerate(self.info['modulos'], 1):
            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)

            # Slide 1: Título del módulo
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = f"Módulo {i}: {modulo['nombre']}"
            subtitle.text = f"{self.info['nombre']}\n{self.info['grupo_edad']}"

            # Slide 2: Objetivos del módulo
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title = slide.shapes.title
            content = slide.placeholders[1]
            title.text = "Objetivos del Módulo"

            tf = content.text_frame
            tf.text = modulo['descripcion']

            # Slides adicionales según el contenido del módulo
            if 'contenido_detallado' in modulo:
                for tema in modulo['contenido_detallado']:
                    slide = prs.slides.add_slide(prs.slide_layouts[1])
                    title = slide.shapes.title
                    content = slide.placeholders[1]
                    title.text = tema['titulo']

                    tf = content.text_frame
                    for punto in tema['puntos']:
                        p = tf.add_paragraph()
                        p.text = punto
                        p.level = 0

            # Slide final: Resumen
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title = slide.shapes.title
            title.text = "Resumen y Próximos Pasos"

            archivo = os.path.join(carpeta_ppts, f"Modulo_{i}_{modulo['nombre'].replace(' ', '_')}.pptx")
            prs.save(archivo)


def main():
    """Función principal para ejecutar generación de cursos"""
    # Este archivo será importado por los scripts de cada curso
    pass


if __name__ == "__main__":
    main()
